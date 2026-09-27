"""Pull text out of a file the extension sends as base64.

PDF and DOCX are parsed locally. Images go to Gemini (multimodal) so Connie can
catch a photographed ID card or a screenshot of a bank statement.
"""
import base64
import io

from docx import Document
from pypdf import PdfReader

from connie import gemini

MAX_BYTES = 10 * 1024 * 1024
MAX_CHARS = 50_000
IMAGE_TYPES = {"image/png", "image/jpeg", "image/webp", "image/heic", "image/heif"}


def extract(file_b64: str, file_name: str, mime: str) -> tuple[str, str]:
    """Returns (text, method). method is 'pdf', 'docx', 'gemini-vision', or 'text'."""
    data = base64.b64decode(file_b64)
    if len(data) > MAX_BYTES:
        raise ValueError("file too large")
    name = (file_name or "").lower()

    if mime == "application/pdf" or name.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(data))
        text = "\n".join(page.extract_text() or "" for page in reader.pages[:30])
        return text[:MAX_CHARS], "pdf"

    if name.endswith(".docx"):
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        parts += [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
        return "\n".join(parts)[:MAX_CHARS], "docx"

    if mime in IMAGE_TYPES:
        return gemini.read_image(data, mime)[:MAX_CHARS], "gemini-vision"

    return data.decode("utf-8", errors="replace")[:MAX_CHARS], "text"
